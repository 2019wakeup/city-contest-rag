"""Final content spot-check of the workbook and the report."""

import io
import json

import openpyxl

XLSX = r"D:\code\City_Contest\文献调研.xlsx"
MD = r"D:\code\City_Contest\文献调研报告.md"

wb = openpyxl.load_workbook(XLSX)
ws = wb["全部文献"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
hdr = [c.value for c in wb["全部文献"][1]]
idx = {h: i for i, h in enumerate(hdr)}

print("sheets:", len(wb.sheetnames))
print("main rows:", len(rows))
print("non-empty IF:", sum(1 for r in rows if str(r[idx["IF"]] or "").strip()))
print("non-empty 引用率:", sum(1 for r in rows if str(r[idx["引用率"]] or "").strip()))
print("rows with code link:", sum(1 for r in rows if str(r[idx["代码链接"]] or "").strip()))
print("rows with dataset url:", sum(1 for r in rows if str(r[idx["数据集下载地址"]] or "").strip()))
print("rows with DOI:", sum(1 for r in rows if str(r[idx["DOI"]] or "").strip()))
print("rows missing 任务（中文）:", sum(1 for r in rows if not str(r[idx["任务（中文）"]] or "").strip()))
print("rows missing 核验方式:", sum(1 for r in rows if not str(r[idx["核验方式"]] or "").strip()))
print("distinct 期刊/会议名:", len({str(r[idx["期刊/会议名"]]) for r in rows}))

print("\n--- sample: 任务2 材料-性能关系 ---")
n = 0
for r in rows:
    if r[idx["类型"]] == "材料-性能关系":
        print("*", str(r[idx["论文标题"]])[:78])
        print("   ", str(r[idx["任务（中文）"]])[:120])
        print("    DOI:", r[idx["DOI"]], "| 核验:", r[idx["核验方式"]], "| IF:", r[idx["IF"]])
        n += 1
        if n >= 3:
            break

print("\n--- sample: rows carrying verified code links ---")
n = 0
for r in rows:
    if str(r[idx["代码链接"]] or "").startswith("http") and r[idx["类型"]] == "评估方法":
        print("*", str(r[idx["论文标题"]])[:70], "->", r[idx["代码链接"]])
        n += 1
        if n >= 4:
            break

txt = io.open(MD, encoding="utf-8").read()
print("\nreport chars:", len(txt), "| lines:", txt.count("\n") + 1)
print("report mentions 549.04:", "549.04" in txt)
print("report mentions JTG 3410-2025:", "3410-2025" in txt)
print("report has appendix A/B/C:",
      all(s in txt for s in ("## 附录 A", "## 附录 B", "## 附录 C")))
print("report left unfilled placeholders:", txt.count("<!--N_"))
bad = [s for s in ("FATAL", "TODO", "TBD", "XXX", "Lorem") if s in txt]
print("suspicious tokens:", bad)
