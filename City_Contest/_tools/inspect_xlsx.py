import json
import sys

import openpyxl

out = {}
for path in [r"D:\code\City_Contest\20261012-分工表.xlsx", r"D:\code\City_Contest\文献调研.xlsx"]:
    wb = openpyxl.load_workbook(path, data_only=False)
    wb_info = {}
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            vals = [("" if c.value is None else str(c.value)) for c in r]
            if any(v.strip() for v in vals):
                rows.append({"row": r[0].row, "cells": vals})
        wb_info[ws.title] = {
            "dims": ws.dimensions,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merged": [str(m) for m in ws.merged_cells.ranges],
            "rows": rows,
        }
    out[path] = wb_info

with open(r"D:\code\City_Contest\_tools\workbook_dump.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)

print("ok")
